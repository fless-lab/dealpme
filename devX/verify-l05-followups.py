"""Vérification du suivi V1-110/111/117–119 et archivage des preuves de recette.

Exécuter après build_suivi.py, puis avec --recalculated après recalcul LibreOffice
du classeur et du scénario L05_100pct_bloque.xlsx dans /tmp/opencode/recalc-dealpme.
Le commit de comparaison est stable ; --base HEAD vérifie la livraison avant commit.
"""
import argparse
import hashlib
import io
import json
import subprocess
from pathlib import Path

import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument("--base", default="73f5621")
parser.add_argument("--recalculated", action="store_true")
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]


def git(*arguments):
    return subprocess.check_output(["git", *arguments], cwd=root)


def tasks(book):
    rows = [r for r in book["Taches"].iter_rows(min_row=6, values_only=True) if r[1]]
    result = {r[1]: r for r in rows}
    assert len(result) == len(rows), "Identifiants dupliqués"
    return result


before = openpyxl.load_workbook(io.BytesIO(git("show", f"{args.base}:DealPME_Suivi.xlsx")))
after = openpyxl.load_workbook(root / "DealPME_Suivi.xlsx")
b, a = tasks(before), tasks(after)
added = {"V1-117", "V1-118", "V1-119"}
completed = added | {"V1-110", "V1-111"}
assert len(b) == 266 and len(a) == 269 and a.keys() - b.keys() == added and b.keys() <= a.keys()
for tid, old in b.items():
    for index, value in enumerate(old):
        if tid in {"V1-110", "V1-111"} and index in {9, 10, 17, 21}:
            continue
        if tid in {"V1-007", "V1-105"} and index == 21:
            continue
        assert a[tid][index] == value, (tid, index, value, a[tid][index])
assert sum(r[9] == "Complétée" for r in a.values()) == 87
for tid in completed:
    row = a[tid]
    assert row[9] == "Complétée" and row[10] == 1 and row[15] <= row[17]
    assert row[14] == "Abdou-Raouf" and row[18] and len(row[6]) > 100
    assert "2026-09-03" <= str(row[15].date()) <= str(row[16].date()) <= "2026-10-15"
for tid in {"V1-112", "V1-113", "V1-114", "V1-115", "V1-116"}:
    assert a[tid][10] == 0 and a[tid][17] is None
assert sum(a[tid][7] for tid in added) == 5
for version, total in [("V1", 350), ("V2", 228), ("V3", 219), ("V4", 111), ("V5", 153)]:
    assert sum(r[7] for r in a.values() if r[2] == version) == total
for r in range(6, 11):
    for c in (2, 4, 5):
        assert before["Versions"].cell(r, c).value == after["Versions"].cell(r, c).value
lots = {r[1]: r for r in after["Plan_execution"].iter_rows(min_row=6, values_only=True) if r[1]}
for tid in added:
    assert [lot for lot, row in lots.items() if tid in str(row[4]).split(", ")] == ["L05"]
assert "COUNTIFS" in lots["L05"][11]
assert (root / "DealPME_Suivi.ods").read_bytes() == git("show", f"{args.base}:DealPME_Suivi.ods")
print("PASS : 266 anciennes tâches préservées hors clôtures/commentaires autorisés ; 269 IDs uniques, 87 complétées, +5 j/p, fenêtres et archive ODS inchangées.")

folder = Path("/tmp/opencode/l05-tracking-source")
if not args.recalculated:
    ids = set(lots["L05"][4].split(", "))
    for row in after["Taches"].iter_rows(min_row=6):
        if row[1].value in ids:
            row[9].value = "Bloquée" if row[1].value == "V1-112" else "Complétée"
            row[10].value = 1
    folder.mkdir(exist_ok=True)
    after.save(folder / "L05_100pct_bloque.xlsx")
else:
    recalculated = Path("/tmp/opencode/recalc-dealpme")
    for filename in ["DealPME_Suivi.xlsx", "L05_100pct_bloque.xlsx"]:
        book = openpyxl.load_workbook(recalculated / filename, data_only=True)
        assert not [(s.title, c.coordinate, c.value) for s in book for row in s for c in row if c.data_type == "e"]
        lot = next(r for r in book["Plan_execution"].iter_rows(values_only=True) if r[1] == "L05")
        assert lot[11] == "Amorcé"
        if filename.startswith("L05_"):
            assert lot[10] == 1
        else:
            assert 0 < lot[10] < 1
            recalced = tasks(book)
            for tid in {"V1-112", "V1-113", "V1-114", "V1-115", "V1-116"}:
                assert recalced[tid][11] == 0
    print("PASS : formules recalculées sans erreur ; L05 reste ouvert même à 100 % avec une recette bloquée.")

    def load(path):
        return json.loads((root / path).read_text())

    reports = {}
    for filename in ["smoke-results.json", "gate-rejections.json", "l02-results.json", "l03-results.json", "l04-registry-results.json", "l04-alert-results.json", "l04-operations-results.json", "l04-restore-results.json", "l05-results.json", "l05-capture-results.json", "remo-api-results.json"]:
        report = load(f".ci-artifacts/{filename}")
        if filename == "l05-capture-results.json":
            assert report["synthetic"] is True and report["bytes"] > 0 and report["stopReason"] == "admission_revoked"
        else:
            assert report["status"] == "PASS", filename
        for check in report.get("checks", []):
            assert check.get("status") == "PASS" or check.get("passed") is True, (filename, check)
        reports[filename] = report
    assert reports["smoke-results.json"]["scope"] == "all"
    matrix = load("qa/l05-wiring-matrix.json")
    assert matrix["evidence"] == "FULL_ISOLATED_SUITE_PASS" and matrix["routeCount"] == 27
    for filename, evidence in matrix["reports"].items():
        assert hashlib.sha256((root / ".ci-artifacts" / filename).read_bytes()).hexdigest() == evidence["sha256"]
    findings = load("qa/l05-wiring-findings.json")
    assert all(f["status"] == "CLOSED" for f in findings["findings"])
    paths = set(git("diff", "--name-only", args.base).decode().splitlines())
    paths.update(git("ls-files", "--others", "--exclude-standard").decode().splitlines())
    source_hashes = {p: hashlib.sha256((root / p).read_bytes()).hexdigest() for p in sorted(paths) if (root / p).is_file() and (p.startswith(("codebases/", "packages/", "devX/", ".github/")) or p in {"package.json", "qa/registre-interactions.json"})}
    result = {
        "status": "PASS", "sourceBase": git("rev-parse", args.base).decode().strip(),
        "verifiedAt": reports["smoke-results.json"]["finishedAt"],
        "scope": "V1-110/111/117/118/119 hors compte fournisseur",
        "providerQualified": False, "sourceHashes": source_hashes,
        "checksExecutedInSession": {"builds": 26, "unitTests": 175, "unitFiles": 18, "typecheck": "PASS", "lint": "PASS", "auditHighCritical": 0, "auditModerate": 4},
        "tracking": {"previousTasks": 266, "tasks": 269, "completed": 87, "newIds": sorted(added), "closedIds": sorted(completed), "addedPersonDays": 5, "personDays": 1061, "versionWindowsPreserved": True, "odsUnchanged": True, "recalculatedWithoutErrors": True, "l05StillOpenAt100PercentIfBlocked": True},
        "reports": reports,
        "openFollowups": ["V1-112", "V1-113", "V1-114", "V1-115", "V1-116", "V1-107"],
        "verificationCommands": ["npm ci && npm run build && npm run typecheck && npm test && npm run lint && npm audit --audit-level=high", "npm run ci:verify-gates", "PLAYWRIGHT_BROWSERS_PATH=/tmp/opencode/dealpme-browsers npm run ci:smoke", "npm run ci:l05-wiring -- --evidence"],
    }
    (root / "qa/l05-wiring-verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print("PASS : preuves de régression et empreintes archivées dans qa/l05-wiring-verification.json.")
